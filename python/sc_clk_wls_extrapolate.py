#!/usr/bin/env python3
"""sc_clk_wls_extrapolate.py - Build a -clkfix clock file for the drone (SC)
side from a GS-recomputed AOWR clock_diff_s/tag_tow series (e.g. a GS
receiver clock offset recomputed with IGS Rapid precise ephemeris), by
faithfully replicating the onboard SC AOWR WLS recurrence in sdr_pvt.c
(sdr_pvt_udobs()'s SC-mode block, ~line 1757-1819):

    new_clock = s_sc_last_dt_aowr + clock_diff_s
    new_tow   = wrap(tag_tow + s_sc_last_dt_aowr, [0, 604800))
    <push (new_tow, new_clock) into a 3-sample history, WLS-fit
     y = dt_i + drift*(t - t0), then extrapolate to each GNSS obs epoch>

s_sc_last_dt_aowr is NOT the small (~ms) final clock bias -- on this
project's PS test setup it is a large (~4.3e5 s) constant, because the raw
PS pseudorange itself carries a large injected offset (see aowr.csv's
dt_raw/dt_cp fields, ~432198.7 s for a converged pass). It comes from the
SC's own onboard PS-pass processing ($AOWR log / aowr.csv's dt_cp column,
held at its last converged value -- sdr_pvt.c:1676-1678), NOT from the GS
file. clock_diff_s alone is not a usable clock bias; only
(seed + clock_diff_s) is.

Input xlsx columns used: 'rx TOW' (tag_tow) and 'dt_aowr-gnss_sc'
(clock_diff_s); other columns in that sheet are the GS-side quantities and
an already-differenced check value, not used here.

--gs-csv (recommended): the xlsx's 'dt_aowr-gnss_sc' column is itself an
Excel FORMULA, not raw data -- '=B+H+Sheet1!$B$2-Sheet1!$B$1', i.e.
dt_aowr-gnss_gs + dt_total + (dt_sc_s - dt_gs_s) (reverse-engineered by
reading the xlsx's stored formula strings, not guessed). Column B
(dt_aowr-gnss_gs, ~-4.3e5 s magnitude) is the dominant term, and its value
as stored in the xlsx is only accurate to ~1e-9 s -- a full IEEE-754 double
has ~15-17 significant decimal digits total, and ~432198.xxx already spends
6 of those left of the decimal, so whatever produced the xlsx (a round-trip
through Excel's own display/copy-paste) left only ~1e-9 s of fractional
precision, capping dt_aowr-gnss_sc at that same absolute precision even
though the small quantities (H/dt_total, ~0.15 s or less) it's added to are
independently exact. --gs-csv points at a plain-text re-export of column B
at ~1e-10 s precision (e.g. dt_aowr_gnss_sc_RapidPrecise.csv, 2 columns: rx
TOW, dt_aowr-gnss_gs, no header); when given, dt_aowr-gnss_sc is recomputed
from that precise B plus H/Sheet1 constants read straight out of the xlsx
(H is a simple small+small addition internally, no precision loss there) --
verified this reproduces the xlsx's own cached dt_aowr-gnss_sc values to
within ~3e-9 s (recovering the ~1e-9 s of precision the round-trip cost).

--min-elapsed-s (staged clock-fix): don't emit -clkfix entries until this
many seconds after rover.obs's first epoch. Investigation on this session
found the epoch-loss from clock-fixing wasn't from clock noise or bias --
it was a self-reinforcing cascade from switching on the tight external
clock constraint (clock_bias_fixed) *before* the PPP KF's own ambiguities
had converged from cold start (nconv, the # of converged double-difference
ambiguities, was still climbing from 0 through minute ~10 of this session,
while clock-fixing started at minute ~6.7): an immature epoch's residual
can't be absorbed by the pinned clock, a satellite gets rejected, its
ambiguity resets, nconv drops further, and the loop feeds itself. Delaying
clock-fix start until convergence (here, minute 10 -> --min-elapsed-s 600)
let the filter free-run through the fragile cold-start stretch, then
recovered ~100% of free-running's epoch count for the rest of the session
with the clock prior fully unchanged (8 m, not loosened). This is a
time-based proxy for "has nconv converged", tuned to this one session's
convergence curve -- it is NOT a general convergence detector, and a
different flight/pass would need its own value (or a real nconv-gated
version ported into sdr_pvt.c/postpos.c, which this flags as a gap in the
onboard clock_bias_fixed logic: it currently engages the moment an AOWR
clock is available, with no check on the KF's own convergence state).

Usage:
    python3 sc_clk_wls_extrapolate.py work/20260710101607 \\
        --xlsx dt_aowr_gnss_sc_RapidPrecise.xlsx \\
        --gs-csv dt_aowr_gnss_sc_RapidPrecise.csv \\
        --aowr aowr.csv --obs rover.obs --min-elapsed-s 600 \\
        -o clk_sc_precise_new.txt
"""
import sys, os, argparse
from datetime import datetime, timedelta
import numpy as np
import pandas as pd

SC_CLOCK_HIST_SIZE = 3     # mirrors sdr_pvt.c's SC_CLOCK_HIST_SIZE
WEEK_S = 604800.0
GPST_EPOCH = datetime(1980, 1, 6)  # GPST calendar, no leap-second correction


def datetime_to_tow(dt):
    total = (dt - GPST_EPOCH).total_seconds()
    week = int(total // WEEK_S)
    return week, total - week * WEEK_S


def tow_to_datetime(week, tow):
    return GPST_EPOCH + timedelta(days=week * 7, seconds=tow)


def load_dt_cp_seed(aowr_csv):
    """Converged onboard dt_aowr_cp == s_sc_last_dt_aowr for this session
    (sdr_pvt.c:1676-1678: held-last-value once count>0 and dt_cp!=0)."""
    df = pd.read_csv(aowr_csv)
    conv = df[(df['count'] > 0) & (df['dt_cp'] != 0)]
    if conv.empty:
        raise SystemExit(f'{aowr_csv}: no converged (count>0, dt_cp!=0) rows -- '
                          'cannot determine s_sc_last_dt_aowr seed')
    seed = float(conv['dt_cp'].iloc[-1])
    spread = conv['dt_cp'].max() - conv['dt_cp'].min()
    print(f'seed dt_cp (s_sc_last_dt_aowr) = {seed!r} '
          f'(held over {len(conv)} rows, spread {spread:.3e} s)')
    return seed


def wls_fit(tow_hist, clk_hist):
    """y = dt_i + drift*(t - t0) -- mirrors sdr_pvt.c:1788-1804."""
    t0 = tow_hist[0]
    n = len(tow_hist)
    if n <= 1:
        return clk_hist[-1], 0.0, t0
    t = np.array(tow_hist) - t0
    y = np.array(clk_hist)
    S0, S1, S2 = float(n), t.sum(), (t * t).sum()
    T0, T1 = y.sum(), (t * y).sum()
    det = S0 * S2 - S1 * S1
    if abs(det) < 1e-12:
        return clk_hist[-1], 0.0, t0
    dt_i = (S2 * T0 - S1 * T1) / det
    drift = (-S1 * T0 + S0 * T1) / det
    return dt_i, drift, t0


def load_clock_diff_xlsx(xlsx_path, sheet):
    """Plain read of the xlsx's own (Excel-computed) dt_aowr-gnss_sc column --
    subject to the ~1e-9 s precision loss described in the module docstring.
    Returns a ['rx TOW', 'dt_aowr-gnss_sc'] DataFrame, sorted, NaNs dropped."""
    df = pd.read_excel(xlsx_path, sheet_name=sheet)
    df = df[['rx TOW', 'dt_aowr-gnss_sc']].dropna().sort_values('rx TOW')
    if df.empty:
        raise SystemExit(f'{xlsx_path}: no valid (rx TOW, dt_aowr-gnss_sc) rows')
    return df


def load_clock_diff_precise(xlsx_path, sheet, gs_csv_path):
    """Recompute dt_aowr-gnss_sc = dt_aowr-gnss_gs + dt_total +
    (dt_sc_s - dt_gs_s) using a full-precision (~1e-10 s) re-export of
    dt_aowr-gnss_gs (gs_csv_path: 2 plain columns, rx TOW and
    dt_aowr-gnss_gs, no header) instead of the xlsx's own rounded copy of
    that column. dt_total (H) and the Sheet1 dt_gs_s/dt_sc_s constants are
    read straight from the xlsx's cached formula results -- those are
    small+small additions internally, no precision loss to fix there.
    Returns a ['rx TOW', 'dt_aowr-gnss_sc'] DataFrame, sorted, NaNs dropped."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet]
    dt_gs_s = wb['Sheet1']['B1'].value
    dt_sc_s = wb['Sheet1']['B2'].value
    const = dt_sc_s - dt_gs_s

    rows = []
    for r in range(2, ws.max_row + 1):
        tow = ws.cell(row=r, column=1).value   # A: rx TOW
        h = ws.cell(row=r, column=8).value     # H: dt_total (cached)
        if tow is None or h is None:
            continue
        rows.append((tow, h))
    h_df = pd.DataFrame(rows, columns=['rx TOW', 'dt_total'])

    gs_df = pd.read_csv(gs_csv_path, header=None, names=['rx TOW', 'dt_aowr-gnss_gs'])

    df = h_df.merge(gs_df, on='rx TOW', how='inner')
    if df.empty:
        raise SystemExit(f'{gs_csv_path}: no rows with matching rx TOW in {xlsx_path}')
    df['dt_aowr-gnss_sc'] = df['dt_aowr-gnss_gs'] + df['dt_total'] + const
    return df[['rx TOW', 'dt_aowr-gnss_sc']].sort_values('rx TOW')


def build_updates(df, seed_dt_cp, offset_s=0.0, hist_size=SC_CLOCK_HIST_SIZE):
    """Replicate sdr_pvt.c:1757-1819's per-AOWR-sample update. Returns a
    list of (new_tow, dt_i, drift, t0) WLS regression states in time order,
    each valid from its new_tow until superseded by the next update.

    df: ['rx TOW', 'dt_aowr-gnss_sc'] as returned by load_clock_diff_xlsx()
    or load_clock_diff_precise().
    offset_s: constant added to every new_clock sample before the WLS fit --
    for testing a hypothesized uncalibrated GS antenna-cable delay (does not
    affect new_tow, which is derived from tag_tow + seed only).
    hist_size: WLS regression window (# of AOWR samples), default matches
    the onboard SC_CLOCK_HIST_SIZE=3 -- larger smooths more (lower variance)
    but responds to real drift-rate changes more slowly (see --hist-size)."""
    tow_hist, clk_hist, updates = [], [], []
    for tag_tow, clock_diff_s in df.itertuples(index=False):
        new_clock = seed_dt_cp + clock_diff_s + offset_s
        new_tow = (tag_tow + seed_dt_cp) % WEEK_S

        if len(tow_hist) < hist_size:
            tow_hist.append(new_tow); clk_hist.append(new_clock)
        else:
            tow_hist, clk_hist = tow_hist[1:] + [new_tow], clk_hist[1:] + [new_clock]

        dt_i, drift, t0 = wls_fit(tow_hist, clk_hist)
        updates.append((new_tow, dt_i, drift, t0))
    return updates


def extrapolate(updates, obs_tows):
    """clock_est at each obs epoch tow -- mirrors sdr_pvt.c:1807-1819: use
    whichever WLS regression state was most recently established at/before
    this epoch. Epochs before the first AOWR update are skipped (no entry
    -> rnx2rtkp falls back to normal free-clock estimation for them, same
    as onboard's sc_clk_ready==0 gate)."""
    out, j = [], -1
    for ot in obs_tows:
        while j + 1 < len(updates) and updates[j + 1][0] <= ot:
            j += 1
        if j < 0:
            continue
        _, dt_i, drift, t0 = updates[j]
        dt = ot - t0
        if dt < 0:
            dt += WEEK_S
        out.append((ot, dt_i + drift * dt))
    return out


def obs_epoch_range(obs_path):
    """First/last '>' epoch datetime in a RINEX obs file (GPST calendar)."""
    first = last = None
    with open(obs_path) as f:
        for line in f:
            if not line.startswith('>'):
                continue
            p = line.split()
            y, mo, d, h, mi = (int(p[1]), int(p[2]), int(p[3]), int(p[4]), int(p[5]))
            s = float(p[6])
            dt = datetime(y, mo, d, h, mi) + timedelta(seconds=s)
            if first is None:
                first = dt
            last = dt
    if first is None:
        raise SystemExit(f'{obs_path}: no epoch (">") lines found')
    return first, last


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                  formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('session_dir')
    ap.add_argument('--xlsx', default='dt_aowr_gnss_sc_RapidPrecise.xlsx')
    ap.add_argument('--sheet', default='dt_aowr_gnss_sc_RapidPrecise')
    ap.add_argument('--aowr', default='aowr.csv')
    ap.add_argument('--obs', default='rover.obs')
    ap.add_argument('--gs-csv', default=None,
                    help='full-precision (~1e-10 s) plain re-export of the '
                         'dt_aowr-gnss_gs column (2 cols: rx TOW, value, no '
                         'header) -- when given, dt_aowr-gnss_sc is recomputed '
                         'from this instead of read from the xlsx directly '
                         '(see module docstring: the xlsx column is only good '
                         'to ~1e-9 s due to an Excel round-trip)')
    ap.add_argument('-o', '--out', default='clk_sc_precise_new.txt')
    ap.add_argument('--offset-m', type=float, default=0.0,
                    help='constant offset (metres, distance domain) added to '
                         'every clock sample before WLS -- e.g. +8/-8 to test '
                         'a hypothesized uncalibrated GS antenna-cable delay')
    ap.add_argument('--min-elapsed-s', type=float, default=0.0,
                    help='staged clock-fix: suppress -clkfix entries until '
                         'this many seconds after rover.obs\'s first epoch, '
                         'to keep clock_bias_fixed off during PPP KF cold-'
                         'start convergence -- see module docstring. 0 (default) '
                         'disables staging (clkfix starts as soon as AOWR data '
                         'does, the old behaviour)')
    ap.add_argument('--hist-size', type=int, default=SC_CLOCK_HIST_SIZE,
                    help=f'WLS regression window, # of AOWR samples (default '
                         f'{SC_CLOCK_HIST_SIZE}, matching onboard '
                         'SC_CLOCK_HIST_SIZE) -- larger smooths more but '
                         'responds to real drift-rate changes more slowly')
    args = ap.parse_args()

    sd = args.session_dir.rstrip('/')
    xlsx_path = os.path.join(sd, args.xlsx)
    aowr_path = os.path.join(sd, args.aowr)
    obs_path = os.path.join(sd, args.obs)
    out_path = os.path.join(sd, args.out)
    offset_s = args.offset_m / 299792458.0
    if offset_s:
        print(f'applying offset {args.offset_m:+.3f} m = {offset_s*1e9:+.3f} ns to every sample')

    if args.gs_csv:
        gs_csv_path = os.path.join(sd, args.gs_csv)
        clk_df = load_clock_diff_precise(xlsx_path, args.sheet, gs_csv_path)
        print(f'using precise dt_aowr-gnss_sc recomputed from {args.gs_csv} '
              f'({len(clk_df)} rows)')
    else:
        clk_df = load_clock_diff_xlsx(xlsx_path, args.sheet)

    seed = load_dt_cp_seed(aowr_path)
    updates = build_updates(clk_df, seed, offset_s, args.hist_size)
    print(f'{len(updates)} AOWR update(s); first new_tow={updates[0][0]:.3f} '
          f'last new_tow={updates[-1][0]:.3f}')

    first_dt, last_dt = obs_epoch_range(obs_path)
    week, first_tow = datetime_to_tow(first_dt)
    _, last_tow = datetime_to_tow(last_dt)
    print(f'rover.obs epochs: {first_dt} .. {last_dt} (GPST week {week}, '
          f'tow {first_tow:.1f}..{last_tow:.1f})')

    n = int(round(last_tow - first_tow)) + 1
    obs_tows = [first_tow + k for k in range(n)]
    rows = extrapolate(updates, obs_tows)

    if args.min_elapsed_s > 0:
        gate_tow = first_tow + args.min_elapsed_s
        n_before = len(rows)
        rows = [(tow, dtr) for tow, dtr in rows if tow >= gate_tow]
        print(f'staged: suppressed {n_before - len(rows)} entries before '
              f'{args.min_elapsed_s:.0f}s elapsed (gate_tow={gate_tow:.1f}) '
              '-- filter free-runs through cold-start convergence')

    with open(out_path, 'w') as f:
        f.write('# time,dtr_s -- generated by sc_clk_wls_extrapolate.py from '
                f'{args.xlsx}\n')
        for tow, dtr in rows:
            dt = tow_to_datetime(week, tow)
            f.write(f"{dt.strftime('%Y/%m/%d %H:%M:%S')}.{int(dt.microsecond/1000):03d} "
                    f"{dtr:.12f}\n")

    print(f'wrote {len(rows)} entries to {out_path} '
          f'({rows[0][1]:.9f} .. {rows[-1][1]:.9f} s)' if rows else
          f'wrote 0 entries to {out_path} (no obs epoch reached after the first AOWR update)')


if __name__ == '__main__':
    main()
